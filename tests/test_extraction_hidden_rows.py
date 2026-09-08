"""
Una fila oculta no se lee.

Las Traffic Sheets traen ejemplos de la plantilla ocultos entre el
trabajo real ("client_geo_creativename_728x90" y compania). Leerlos
metia en el QA creativos que nadie pidio revisar: los hallazgos de la
solicitud real quedaban enterrados bajo los del ejemplo. La regla del
equipo es corta: lo oculto no se lee.

Run with pytest, or directly:
    python tests/test_extraction_hidden_rows.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402

from core.extraction import read_sheet  # noqa: E402


def _sheet(rows: list[list[object]], hidden: set[int],
           merged: str | None = None):
    """Escribe las filas en un .xlsx temporal y devuelve (grid, anomalies)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in rows:
        ws.append(row)
    for index in hidden:
        ws.row_dimensions[index].hidden = True
    if merged:
        ws.merge_cells(merged)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "ts.xlsx"
        wb.save(path)
        return read_sheet(path, "Data")


def _texts(grid) -> list[str]:
    return [
        cell.text
        for row in grid.rows.values()
        for cell in row.values()
        if cell.text
    ]


def test_a_hidden_row_is_not_read():
    grid, _ = _sheet(
        [["Creative Name"], ["real_banner_300x250"],
         ["client_geo_creativename_728x90"]],
        hidden={3},
    )
    assert "real_banner_300x250" in _texts(grid)
    assert "client_geo_creativename_728x90" not in _texts(grid)
    assert 3 not in grid.rows


def test_visible_rows_below_a_hidden_block_survive():
    # Los ejemplos suelen ir en medio, no al final: saltarlos no puede
    # cortar la lectura de lo que viene despues.
    rows = [["Creative Name"], ["real_a"]]
    rows += [[f"example_{i}"] for i in range(3, 20)]
    rows += [["real_b"]]
    grid, _ = _sheet(rows, hidden=set(range(3, 20)))
    texts = _texts(grid)
    assert "real_a" in texts
    assert "real_b" in texts
    assert not [t for t in texts if t.startswith("example_")]


def test_max_row_ignores_the_hidden_tail():
    # max_row alimenta el recorrido de las reglas. Si contara filas
    # ocultas, QA2 volveria a mirarlas por otra puerta.
    grid, _ = _sheet(
        [["Creative Name"], ["real_a"], ["example_1"], ["example_2"]],
        hidden={3, 4},
    )
    assert grid.max_row == 2


def test_hidden_rows_are_reported_not_silent():
    # Saltar filas cambia lo que el QA revisa: tiene que quedar dicho.
    _, anomalies = _sheet(
        [["Creative Name"], ["real_a"], ["example_1"]],
        hidden={3},
    )
    hidden = [a for a in anomalies if a.code == "EXT-HIDDEN-ROWS"]
    assert len(hidden) == 1
    assert hidden[0].detail["rows"] == [3]


def test_a_sheet_without_hidden_rows_says_nothing():
    grid, anomalies = _sheet(
        [["Creative Name"], ["real_a"], ["real_b"]],
        hidden=set(),
    )
    assert grid.max_row == 3
    assert not [a for a in anomalies if a.code == "EXT-HIDDEN-ROWS"]


def test_a_merge_anchored_on_a_hidden_row_does_not_resurrect_it():
    # El fill-down de celdas combinadas copia el ancla hacia abajo. Si
    # el ancla esta oculta no hay nada que copiar, y el fill no puede
    # reintroducir el texto que acabamos de descartar.
    grid, _ = _sheet(
        [["Creative Name"], ["real_a"], ["example_1"], [None]],
        hidden={3},
        merged="A3:A4",
    )
    assert "example_1" not in _texts(grid)


if __name__ == "__main__":
    passed = failed = 0
    for name, fn in sorted(globals().items()):
        if not name.startswith("test_") or not callable(fn):
            continue
        try:
            fn()
        except AssertionError as exc:
            failed += 1
            print(f"FAIL {name}: {exc}")
        else:
            passed += 1
            print(f"ok   {name}")

    print(f"\n{passed} passed, {failed} failed")
    sys.exit(1 if failed else 0)
