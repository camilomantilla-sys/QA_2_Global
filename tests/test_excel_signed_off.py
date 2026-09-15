"""
Los tres colores del Excel, y la celda de firma que sobraba.

Verde y naranja decian si los dos lados coinciden. Faltaba el tercer
caso, que es el que mas importa al abrir el archivo: no coincide PERO
una persona lo firmo -- las fechas de creativo que arrancan antes para
testear, casi siempre. En naranja se confundia con lo que sigue
abierto.

Y habia dos sitios donde escribir la misma firma: el desplegable de
aprobacion y una celda de texto libre que pedia editarse a mano. Queda
el desplegable; quien firmo y cuando ya estan en el Implementation
Record.

Run with pytest, or directly:
    python tests/test_excel_signed_off.py
"""
from __future__ import annotations

import dataclasses
import sys
import warnings
from datetime import datetime
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.excel_report import (  # noqa: E402
    AGREE_FILL,
    DISAGREE_FILL,
    SIGNED_OFF_FILL,
    build_excel_report,
)
from core.pdf_report import ReportMeta  # noqa: E402
from core.qa_export import COLUMNS  # noqa: E402


def _meta() -> ReportMeta:
    fields = {}
    for field in dataclasses.fields(ReportMeta):
        if field.default is not dataclasses.MISSING:
            fields[field.name] = field.default
        elif field.default_factory is not dataclasses.MISSING:
            fields[field.name] = field.default_factory()
        else:
            fields[field.name] = ""
    fields.update(
        campaign="Prueba", verdict="NEEDS_REVIEW",
        verdict_label="Review required", ts_campaign_id="313038",
        export_campaign_id="313038", generated_at=datetime.now(),
        metrics={"Worked Placements": 3},
    )
    return ReportMeta(**fields)


def _row(**kw) -> dict:
    row = {column: "" for column in COLUMNS}
    row.update(kw)
    return row


ROWS = [
    _row(Status="PASS", **{
        "TS Creative": "A.png", "Innovid Creative": "A.png",
        "TS Start Date": "2026-09-08", "Innovid Start Date": "2026-09-08"}),
    _row(Status="REVIEW", Notes="INV-001: flights on other dates", **{
        "TS Creative": "B.png", "Innovid Creative": "B.png",
        "TS Start Date": "2026-09-08", "Innovid Start Date": "2026-09-01"}),
    _row(Status="PASS", Notes="MANUALLY Approved by Camilo Mantilla", **{
        "TS Creative": "C.png", "Innovid Creative": "C.png",
        "TS Start Date": "2026-09-08", "Innovid Start Date": "2026-09-01"}),
]


def _workbook():
    blank = pd.DataFrame()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        data = build_excel_report(_meta(), blank, blank, blank, qa_rows=ROWS)
    return load_workbook(BytesIO(data))


def _colour(ws, row: int, column: str) -> str:
    index = {cell.value: cell.column for cell in ws[1]}
    return ws.cell(row, index[column]).fill.start_color.rgb[-6:]


def test_matching_dates_are_green():
    assert _colour(_workbook()["QA"], 2, "TS Start Date") == \
        AGREE_FILL.start_color.rgb[-6:]


def test_an_open_mismatch_is_orange():
    assert _colour(_workbook()["QA"], 3, "TS Start Date") == \
        DISAGREE_FILL.start_color.rgb[-6:]


def test_a_mismatch_signed_off_by_hand_has_its_own_colour():
    assert _colour(_workbook()["QA"], 4, "TS Start Date") == \
        SIGNED_OFF_FILL.start_color.rgb[-6:]


def test_the_three_colours_are_different():
    seen = {
        fill.start_color.rgb
        for fill in (AGREE_FILL, DISAGREE_FILL, SIGNED_OFF_FILL)
    }
    assert len(seen) == 3


def test_both_sides_of_the_pair_are_painted():
    # Pintar solo la columna de Innovid haria leer el color como un
    # veredicto sobre Innovid, no sobre el acuerdo entre los dos.
    ws = _workbook()["QA"]
    for column in ("TS Start Date", "Innovid Start Date"):
        assert _colour(ws, 4, column) == SIGNED_OFF_FILL.start_color.rgb[-6:]


def test_the_legend_explains_the_colours():
    ws = _workbook()["QA"]
    text = " ".join(
        str(cell.value or "")
        for row in ws.iter_rows() for cell in row
    )
    assert "signed this off by hand" in text
    assert "still open" in text


def test_the_free_text_sign_off_cell_is_gone():
    ws = _workbook()["Summary"]
    text = " ".join(
        str(cell.value or "") for row in ws.iter_rows() for cell in row
    )
    assert "edit this cell" not in text
    assert "QA2 Approval: pending" not in text


def test_the_approval_dropdown_stays():
    ws = _workbook()["Summary"]
    values = {
        str(cell.value) for row in ws.iter_rows() for cell in row
    }
    assert "PENDING" in values


if __name__ == "__main__":
    passed = failed = 0
    for name, fn in sorted(globals().items()):
        if not name.startswith("test_") or not callable(fn):
            continue
        try:
            fn()
        except AssertionError as exc:
            failed += 1
            print(f"FAIL {name}: {exc}")
        else:
            passed += 1
            print(f"ok   {name}")

    print(f"\n{passed} passed, {failed} failed")
    sys.exit(1 if failed else 0)


# ── la firma se ve siempre, no solo cuando hay desacuerdo ────────────
#
# Camilo, abriendo el Excel de una corrida real: "el Excel me los
# marcaba las firmas manuales pero solo en la columna de notas y no me
# lo marcaba en morado".
#
# No era una regresion. El morado siempre estuvo atado a que DOS
# columnas llenas no coincidieran, y se firma mucho mas que eso: una
# rotacion que no se pudo comparar deja la columna de Innovid vacia, y
# un NOT_VERIFIED firmado tiene los dos lados iguales. En esos casos la
# firma quedaba solo en Notes -- hay que leer fila por fila para
# encontrarla, que es justo lo que un color evita.

SIGNED = "MANUALLY Approved by Camilo: fechas de prueba"


def _qa_sheet_for(rows: list[dict]):
    empty = pd.DataFrame()
    payload = build_excel_report(_meta(), empty, empty, empty, qa_rows=rows)
    return load_workbook(BytesIO(payload))["QA"]


def _fill(ws, row: int, column: str) -> str:
    header = {c.value: c.column for c in ws[1] if c.value}
    cell = ws.cell(row=row, column=header[column])
    return (cell.fill.start_color.rgb or "") if cell.fill else ""


def _is(fill_rgb: str, fill: object) -> bool:
    return fill_rgb.endswith(fill.start_color.rgb[-6:])


def test_a_signed_row_is_marked_on_status_even_when_the_pair_is_empty():
    """
    El caso que fallaba. Una rotacion que Innovid no trae: no hay nada
    contra que comparar, alguien firma, y la fila salia sin un solo
    color.
    """
    ws = _qa_sheet_for([
        _row(Status="PASS", Notes=SIGNED,
             **{"TS Rotation": "EVEN", "Innovid Rotation": ""}),
    ])
    assert _is(_fill(ws, 2, "Status"), SIGNED_OFF_FILL)


def test_a_signed_row_is_marked_on_status_even_when_both_sides_agree():
    """Un NOT_VERIFIED firmado: los dos lados iguales, firma igual."""
    ws = _qa_sheet_for([
        _row(Status="PASS", Notes=SIGNED,
             **{"TS Rotation": "EVEN", "Innovid Rotation": "EVEN"}),
    ])
    assert _is(_fill(ws, 2, "Status"), SIGNED_OFF_FILL)


def test_a_pair_that_really_agrees_stays_green_on_a_signed_row():
    """
    Lo firmado era otra cosa. Pintar esto de morado diria que aqui
    hubo una discrepancia que no existio.
    """
    ws = _qa_sheet_for([
        _row(Status="PASS", Notes=SIGNED,
             **{"TS Rotation": "EVEN", "Innovid Rotation": "EVEN"}),
    ])
    assert _is(_fill(ws, 2, "Innovid Rotation"), AGREE_FILL)


def test_a_half_empty_pair_on_a_signed_row_is_purple():
    """No se pudo comparar y una persona respondio por ello."""
    ws = _qa_sheet_for([
        _row(Status="PASS", Notes=SIGNED,
             **{"TS Rotation": "EVEN", "Innovid Rotation": ""}),
    ])
    assert _is(_fill(ws, 2, "Innovid Rotation"), SIGNED_OFF_FILL)


def test_a_half_empty_pair_without_a_signature_stays_uncoloured():
    """Sin firma sigue sin color: no hay acuerdo ni desacuerdo."""
    ws = _qa_sheet_for([
        _row(Status="REVIEW",
             **{"TS Rotation": "EVEN", "Innovid Rotation": ""}),
    ])
    assert not _is(_fill(ws, 2, "Innovid Rotation"), SIGNED_OFF_FILL)
    assert not _is(_fill(ws, 2, "Innovid Rotation"), DISAGREE_FILL)


def test_an_empty_pair_on_a_signed_row_is_still_left_alone():
    """Las dos vacias no son una comparacion, firmadas o no."""
    ws = _qa_sheet_for([
        _row(Status="PASS", Notes=SIGNED,
             **{"TS Rotation": "", "Innovid Rotation": ""}),
    ])
    assert not _is(_fill(ws, 2, "Innovid Rotation"), SIGNED_OFF_FILL)


def test_an_unsigned_row_keeps_its_own_status_colour():
    """La firma no puede repintar lo que sigue abierto."""
    ws = _qa_sheet_for([
        _row(Status="FAIL",
             **{"TS Rotation": "EVEN", "Innovid Rotation": "70/30"}),
    ])
    assert not _is(_fill(ws, 2, "Status"), SIGNED_OFF_FILL)
    assert _is(_fill(ws, 2, "Innovid Rotation"), DISAGREE_FILL)


def test_the_legend_no_longer_claims_purple_means_a_mismatch():
    """
    Decia "They differ, and a reviewer signed it off", que ahora seria
    mentira en la celda de Status.
    """
    ws = _qa_sheet_for([_row(Status="PASS", Notes=SIGNED)])
    text = " ".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value
    )
    assert "A reviewer signed this off by hand" in text
