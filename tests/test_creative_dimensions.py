"""
El tamano al que se sirve un creativo, y por que no se veia.

Camilo firmo un mismatch y el Excel le mostro los dos lados iguales:

    "me di cuenta que me cambiaba las dimensiones de un creativo de
     728x90 ... y el otro era 720x50 pero en el export de excel los
     veia iguales eso está mal la idea es firmar pero que se evidencie
     la diferencia"

Tiene razon, y era peor que un color. Las unicas columnas de
dimensiones del entregable -- "TS Dimensions" e "Innovid Dimensions"
-- son las del PLACEMENT, y se repiten identicas en todas sus filas.
Un creativo servido a otro tamano salia con las mismas dos celdas que
uno correcto.

Debajo estaba la causa real: ActualCreative ni siquiera leia la
dimension. El export la trae por fila y se descartaba, asi que el dato
no existia en ninguna parte -- ni para compararlo, ni para mostrarlo.

Firmar sin evidencia es lo contrario de firmar. La firma dice "se que
esto no cuadra y respondo por ello"; si el papel no muestra que no
cuadraba, no queda nada que auditar.

Run with pytest, or directly:
    python tests/test_creative_dimensions.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.matching import (  # noqa: E402
    ActualCreative,
    ActualPlacement,
    CreativeLink,
    ExpectedCreative,
    ExpectedPlacement,
    MatchResult,
    PlacementMatch,
)
from core.qa_export import COLUMNS, PAIRS, build_qa_rows, cells_agree  # noqa: E402


def one_creative(ts_dims: str, innovid_dims: str,
                 placement_dims: str = "728x90") -> dict:
    """La fila del creativo, como sale al entregable."""
    ep = ExpectedPlacement(
        placement_id="111", name="BAN", dims=placement_dims, group_name="G",
    )
    ec = ExpectedCreative(
        name="CRE_v01.jpg", creative_id="900", dims=ts_dims, intent="GREEN",
    )
    ep.creatives.append(ec)

    ap = ActualPlacement(placement_id="111", name="BAN", dims=placement_dims)
    ac = ActualCreative(creative_id="900", name="CRE_v01.jpg", dims=innovid_dims)
    ap.creatives.append(ac)

    pm = PlacementMatch(placement_id="111", expected=ep, actual=ap)
    pm.creative_links.append(CreativeLink(expected=ec, actual=ac))
    return build_qa_rows(MatchResult(matched=[pm]), [])[-1]


# ── el dato existe ───────────────────────────────────────────────────

def test_innovid_creatives_carry_their_own_dimensions():
    """Lo que no se leia. Sin esto no hay nada que comparar."""
    assert "dims" in ActualCreative.__dataclass_fields__


def test_the_export_reads_it():
    from parsers.innovid_export import parse_innovid_export  # noqa: F401
    from core import matching

    source = Path(matching.__file__).read_text(encoding="utf-8")
    assert 'dims=norm_dims(row.values.get("dimensions"))' in source


# ── y se ve ──────────────────────────────────────────────────────────

def test_the_creative_dimensions_are_their_own_columns():
    assert "TS Creative Dims" in COLUMNS
    assert "Innovid Creative Dims" in COLUMNS
    assert ("TS Creative Dims", "Innovid Creative Dims") in PAIRS


def test_the_mismatch_camilo_signed_is_visible():
    """728x90 pedido, 720x50 servido. Lo que salia identico."""
    row = one_creative("728x90", "720x50")
    assert row["TS Creative Dims"] == "728x90"
    assert row["Innovid Creative Dims"] == "720x50"
    assert not cells_agree(
        row["TS Creative Dims"], row["Innovid Creative Dims"]
    )


def test_the_placement_columns_still_say_what_they_said():
    """
    Siguen siendo las del placement, iguales en las dos. Por eso hacian
    falta las otras dos, no para reemplazarlas.
    """
    row = one_creative("728x90", "720x50")
    assert row["TS Dimensions"] == row["Innovid Dimensions"] == "728x90"


def test_a_creative_at_the_right_size_agrees():
    row = one_creative("300x250", "300x250")
    assert cells_agree(
        row["TS Creative Dims"], row["Innovid Creative Dims"]
    )


# ── video: una duracion no es un tamano ──────────────────────────────

def test_a_duration_does_not_invent_a_mismatch():
    """
    La Traffic Sheet escribe "15s" para video. Compararlo contra
    1920x1080 pintaria de naranja todas las filas de video de la
    cuenta -- un desacuerdo que no existe. Se deja el lado de Innovid
    en blanco, que es lo que significa: aqui no hay comparacion.
    """
    row = one_creative("15s", "1920x1080", placement_dims="1920x1080")
    assert row["TS Creative Dims"] == "15s"
    assert row["Innovid Creative Dims"] == ""


def test_an_empty_pair_is_not_coloured():
    from core.qa_export import cells_agree as agree

    row = one_creative("15s", "1920x1080", placement_dims="1920x1080")
    assert not agree(
        row["TS Creative Dims"], row["Innovid Creative Dims"]
    )


# ── la firma deja evidencia ──────────────────────────────────────────

def test_signing_a_mismatch_leaves_the_mismatch_on_the_record():
    """
    La prueba que resume el asunto: despues de firmar, el Excel sigue
    mostrando los dos valores distintos, en morado. Si esto se rompe,
    se vuelve a firmar sobre un papel que no dice nada.
    """
    import dataclasses
    from datetime import datetime
    from io import BytesIO

    import pandas as pd
    from openpyxl import load_workbook

    from core.excel_report import SIGNED_OFF_FILL, build_excel_report
    from core.pdf_report import ReportMeta

    row = one_creative("728x90", "720x50")
    row["Notes"] = "MANUALLY Approved by Camilo: aprobado al tamano servido"

    fields = {
        f.name: (
            f.default if f.default is not dataclasses.MISSING
            else f.default_factory()
            if f.default_factory is not dataclasses.MISSING else ""
        )
        for f in dataclasses.fields(ReportMeta)
    }
    fields.update(campaign="R", verdict="NEEDS_REVIEW", verdict_label="x",
                  generated_at=datetime.now(), metrics={})

    empty = pd.DataFrame()
    payload = build_excel_report(
        ReportMeta(**fields), empty, empty, empty, qa_rows=[row]
    )
    ws = load_workbook(BytesIO(payload))["QA"]
    header = {c.value: c.column for c in ws[1] if c.value}

    def cell(name):
        return ws.cell(row=2, column=header[name])

    assert cell("TS Creative Dims").value == "728x90"
    assert cell("Innovid Creative Dims").value == "720x50"
    for name in ("TS Creative Dims", "Innovid Creative Dims", "Status"):
        rgb = cell(name).fill.start_color.rgb or ""
        assert rgb.endswith(SIGNED_OFF_FILL.start_color.rgb[-6:]), name


if __name__ == "__main__":
    import pytest

    sys.exit(pytest.main([__file__, "-q"]))
