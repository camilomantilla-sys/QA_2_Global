"""
Parser de archivos de tags exportados desde Innovid.

Objetivos:
  - Detectar automáticamente la hoja y el header.
  - Soportar tags 1x1, display y video.
  - Conservar cada tag completo.
  - Extraer URLs incluidas dentro de HTML, JavaScript, iframe, VAST o texto.
  - Detectar Campaign ID, Placement ID y dimensiones dentro del tag.
  - No emitir findings de negocio. Eso corresponde al Rule Engine.
"""
from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from openpyxl import load_workbook

from core.normalize import clean_id, norm_key, norm_text
from core.provenance import Anomaly


warnings.filterwarnings(
    "ignore",
    category=UserWarning,
    module="openpyxl",
)


TAG_1X1_IMPRESSION = "1X1_IMPRESSION"
TAG_1X1_CLICK = "1X1_CLICK"
TAG_DISPLAY_JS = "DISPLAY_JS"
TAG_DISPLAY_IFRAME = "DISPLAY_IFRAME"
TAG_DISPLAY_ASYNC = "DISPLAY_ASYNC"
TAG_DISPLAY_INS = "DISPLAY_INS"
TAG_VIDEO_VAST = "VIDEO_VAST"
TAG_PIXEL = "PIXEL"
TAG_PIXEL_HTML = "PIXEL_HTML"
TAG_UNKNOWN = "UNKNOWN"


CORE_HEADERS = {
    "width",
    "height",
    "type",
    "filename",
    "placementid",
    "thirdpartyid",
    "placementname",
    "startdate",
    "enddate",
    "prismaid",
}


TAG_HEADER_CLASSIFICATION = {
    "ftrack1x1imp": TAG_1X1_IMPRESSION,
    "ftrack1x1click": TAG_1X1_CLICK,
    "staticclicktag1": TAG_1X1_CLICK,
    "updateclicktag1": TAG_1X1_CLICK,
    "jshttps": TAG_DISPLAY_JS,
    "ifhttps": TAG_DISPLAY_IFRAME,
    "asynchttps": TAG_DISPLAY_ASYNC,
    "ins": TAG_DISPLAY_INS,
    "prerollhttps": TAG_VIDEO_VAST,
    "protectedpixel": TAG_PIXEL,
    "protectedpixelhtml": TAG_PIXEL_HTML,
    "protectedpixeljs": TAG_PIXEL,
    "protectedpixeljshtml": TAG_PIXEL_HTML,
    "disqo": TAG_PIXEL,
}


_URL = re.compile(
    r"https?://[^\s\"'<>\\]+",
    re.IGNORECASE,
)

_CAMPAIGN_PATH = re.compile(
    r"(?:click|imp)/\d+/(\d+)[;/]",
    re.IGNORECASE,
)

_PLACEMENT_PATH = re.compile(
    r"(?:click|imp)/\d+/\d+[;,](\d+)[;/]",
    re.IGNORECASE,
)

_PLACEMENT_QUERY = re.compile(
    r"(?:placementId|placement_id|data-placement-id)[=\"':\s]+(\d+)",
    re.IGNORECASE,
)

_WIDTH_QUERY = re.compile(
    r"(?:ft_width|data-placement-width|width)[=\"':\s]+(\d+)",
    re.IGNORECASE,
)

_HEIGHT_QUERY = re.compile(
    r"(?:ft_height|data-placement-height|height)[=\"':\s]+(\d+)",
    re.IGNORECASE,
)

_MACRO = re.compile(
    r"\$\{[^{}]+\}"
    r"|\[[A-Za-z_][A-Za-z0-9_]*\]"
    r"|%%[^%]+%%"
    r"|INSERT_CACHEBUSTER_HERE",
    re.IGNORECASE,
)


@dataclass
class TagValue:
    column_name: str
    tag_type: str
    raw: str
    urls: list[str] = field(default_factory=list)
    hosts: list[str] = field(default_factory=list)
    campaign_ids: list[str] = field(default_factory=list)
    placement_ids: list[str] = field(default_factory=list)
    widths: list[str] = field(default_factory=list)
    heights: list[str] = field(default_factory=list)
    macros: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not bool(self.raw.strip())


@dataclass
class TagRow:
    row: int
    width: str = ""
    height: str = ""
    tag_delivery_type: str = ""
    filename: str = ""
    placement_id: str = ""
    third_party_id: str = ""
    placement_name: str = ""
    start_date: Any = None
    end_date: Any = None
    prisma_id: str = ""
    tags: list[TagValue] = field(default_factory=list)

    @property
    def dimensions(self) -> str:
        if self.width and self.height:
            return f"{self.width}x{self.height}"
        return ""

    @property
    def tag_count(self) -> int:
        return len([tag for tag in self.tags if not tag.is_empty])


@dataclass
class TagsResult:
    path: str = ""
    sheet: str = ""
    metadata: dict[str, str] = field(default_factory=dict)
    header_row: int | None = None
    columns: dict[str, int] = field(default_factory=dict)
    tag_columns: dict[str, int] = field(default_factory=dict)
    rows: list[TagRow] = field(default_factory=list)
    anomalies: list[Anomaly] = field(default_factory=list)

    @property
    def fatal(self) -> bool:
        return any(a.severity == "FATAL" for a in self.anomalies)

    @property
    def campaign_id(self) -> str:
        return self.metadata.get("campaignid", "")

    @property
    def distinct_placements(self) -> int:
        return len({row.placement_id for row in self.rows if row.placement_id})

    @property
    def total_tags(self) -> int:
        return sum(row.tag_count for row in self.rows)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_header(ws, scan_rows: int = 40) -> int | None:
    best_row = None
    best_hits = 0

    for row in range(1, min(ws.max_row, scan_rows) + 1):
        keys = {
            norm_key(_cell_text(ws.cell(row, col).value))
            for col in range(1, ws.max_column + 1)
            if ws.cell(row, col).value is not None
        }

        hits = len(keys & CORE_HEADERS)

        if hits > best_hits:
            best_hits = hits
            best_row = row

    if best_hits >= 5:
        return best_row

    return None


def _select_sheet(wb) -> tuple[str, int | None]:
    candidates: list[tuple[int, str, int]] = []

    for ws in wb.worksheets:
        header_row = _find_header(ws)

        if header_row is None:
            continue

        keys = {
            norm_key(_cell_text(ws.cell(header_row, col).value))
            for col in range(1, ws.max_column + 1)
        }

        score = len(keys & CORE_HEADERS)
        candidates.append((score, ws.title, header_row))

    if not candidates:
        return "", None

    candidates.sort(reverse=True)
    _, sheet_name, header_row = candidates[0]

    return sheet_name, header_row


def _read_metadata(ws, header_row: int) -> dict[str, str]:
    wanted = {
        "mediabuyer",
        "client",
        "campaign",
        "campaignid",
        "department",
    }

    metadata: dict[str, str] = {}

    for row in range(1, header_row):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            key = norm_key(_cell_text(value).rstrip(":"))

            if key not in wanted or key in metadata:
                continue

            for next_col in range(col + 1, ws.max_column + 1):
                candidate = ws.cell(row, next_col).value

                if candidate in (None, ""):
                    continue

                metadata[key] = clean_id(candidate)
                break

    return metadata


def _classify_tag_header(header: str) -> str:
    key = norm_key(header)

    if key in TAG_HEADER_CLASSIFICATION:
        return TAG_HEADER_CLASSIFICATION[key]

    if "click" in key:
        return TAG_1X1_CLICK

    if "preroll" in key or "vast" in key:
        return TAG_VIDEO_VAST

    if "iframe" in key or key.startswith("if"):
        return TAG_DISPLAY_IFRAME

    if "async" in key:
        return TAG_DISPLAY_ASYNC

    if "pixel" in key or "tracker" in key:
        return TAG_PIXEL

    if "tag" in key or "https" in key:
        return TAG_UNKNOWN

    return ""


def _extract_urls(raw: str) -> list[str]:
    output: list[str] = []

    for match in _URL.findall(raw):
        cleaned = (
            match.rstrip(");,")
            .replace("&amp;", "&")
            .replace("\\/", "/")
        )

        if cleaned not in output:
            output.append(cleaned)

    return output


def _extract_hosts(urls: list[str]) -> list[str]:
    output: list[str] = []

    for url in urls:
        try:
            host = urlsplit(url).netloc.casefold()
        except Exception:
            continue

        if host and host not in output:
            output.append(host)

    return output


def _unique_matches(pattern: re.Pattern, raw: str) -> list[str]:
    output: list[str] = []

    for match in pattern.findall(raw):
        value = match if isinstance(match, str) else match[0]

        if value and value not in output:
            output.append(value)

    return output


def _parse_tag_value(column_name: str, raw: str) -> TagValue:
    urls = _extract_urls(raw)

    campaign_ids = _unique_matches(_CAMPAIGN_PATH, raw)

    placement_ids = _unique_matches(_PLACEMENT_PATH, raw)

    for placement_id in _unique_matches(_PLACEMENT_QUERY, raw):
        if placement_id not in placement_ids:
            placement_ids.append(placement_id)

    return TagValue(
        column_name=column_name,
        tag_type=_classify_tag_header(column_name) or TAG_UNKNOWN,
        raw=raw,
        urls=urls,
        hosts=_extract_hosts(urls),
        campaign_ids=campaign_ids,
        placement_ids=placement_ids,
        widths=_unique_matches(_WIDTH_QUERY, raw),
        heights=_unique_matches(_HEIGHT_QUERY, raw),
        macros=sorted(set(_MACRO.findall(raw))),
    )


def parse_innovid_tags(path: Path) -> TagsResult:
    result = TagsResult(path=str(path))

    try:
        wb = load_workbook(
            path,
            data_only=True,
            keep_vba=False,
            read_only=False,
        )
    except Exception as error:
        result.anomalies.append(
            Anomaly(
                "TAG-LOAD-FAILED",
                "FATAL",
                f"No se pudo abrir el archivo de tags: {error}",
            )
        )
        return result

    try:
        sheet_name, header_row = _select_sheet(wb)

        if not sheet_name or header_row is None:
            result.anomalies.append(
                Anomaly(
                    "TAG-HEADER-NOT-FOUND",
                    "FATAL",
                    "No se encontró una hoja con headers de tags reconocibles.",
                )
            )
            return result

        result.sheet = sheet_name
        result.header_row = header_row

        ws = wb[sheet_name]

        result.metadata = _read_metadata(ws, header_row)

        for col in range(1, ws.max_column + 1):
            header = norm_text(_cell_text(ws.cell(header_row, col).value))

            if not header:
                continue

            key = norm_key(header)
            result.columns[key] = col

            tag_type = _classify_tag_header(header)

            if tag_type:
                result.tag_columns[header] = col

        required = {
            "placementid",
            "filename",
            "width",
            "height",
        }

        missing = [
            field_name
            for field_name in required
            if field_name not in result.columns
        ]

        if missing:
            result.anomalies.append(
                Anomaly(
                    "TAG-COLUMN-MISSING",
                    "FATAL",
                    "Columnas requeridas ausentes: "
                    + ", ".join(sorted(missing)),
                )
            )
            return result

        for row_number in range(header_row + 1, ws.max_row + 1):
            placement_cell = ws.cell(
                row_number,
                result.columns["placementid"],
            ).value

            placement_id = clean_id(placement_cell)

            if not placement_id:
                continue

            def value_for(key: str) -> object:
                col = result.columns.get(key)
                return ws.cell(row_number, col).value if col else None

            tag_row = TagRow(
                row=row_number,
                width=clean_id(value_for("width")),
                height=clean_id(value_for("height")),
                tag_delivery_type=norm_text(
                    _cell_text(value_for("type"))
                ),
                filename=norm_text(
                    _cell_text(value_for("filename"))
                ),
                placement_id=placement_id,
                third_party_id=clean_id(
                    value_for("thirdpartyid")
                ),
                placement_name=norm_text(
                    _cell_text(value_for("placementname"))
                ),
                start_date=value_for("startdate"),
                end_date=value_for("enddate"),
                prisma_id=clean_id(value_for("prismaid")),
            )

            for column_name, column_number in result.tag_columns.items():
                raw = _cell_text(
                    ws.cell(row_number, column_number).value
                )

                if not raw:
                    continue

                tag_row.tags.append(
                    _parse_tag_value(column_name, raw)
                )

            result.rows.append(tag_row)

        if not result.rows:
            result.anomalies.append(
                Anomaly(
                    "TAG-NO-DATA",
                    "FATAL",
                    "El archivo no contiene filas con Placement_ID.",
                )
            )

        return result

    finally:
        wb.close()